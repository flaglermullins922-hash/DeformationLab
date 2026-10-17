# -*- coding: utf-8 -*-
"""
md2docx.py —— 把 Markdown 转成 .docx（给用户看的文字交付用）

用法:
    python md2docx.py <输入.md> [输出.docx]

支持: #/##/### 标题、|表格|、```代码块```、- 无序列表、> 引用、
      **加粗**、`行内代码`、--- 分隔线、普通段落。
中文正文/标题统一设 rFonts eastAsia，避免 Word 里中文变成默认西文字体。

依赖: python-docx (系统 python3.11.9 已装)
"""
import sys
import re
import os

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# 字体配置
CN_BODY = "微软雅黑"      # 正文中文字体
CN_HEAD = "微软雅黑"      # 标题中文字体
MONO = "Consolas"         # 代码/行内代码等宽字体
CODE_BG = "F2F2F2"        # 代码块浅灰底纹

BODY_SIZE = 10.5          # 正文五号
H1_SIZE = 20
H2_SIZE = 16
H3_SIZE = 13
H4_SIZE = 12
CODE_SIZE = 9


def set_run_font(run, name=CN_BODY, size=BODY_SIZE, bold=False, color=None, mono=False):
    f = run.font
    f.name = MONO if mono else name
    f.size = Pt(size)
    f.bold = bold
    if color:
        f.color.rgb = RGBColor(*color)
    # 中文字体必须设 eastAsia，否则中文不生效
    r = run._element.rPr
    if r is None:
        r = run._element.get_or_add_rPr()
    rFonts = r.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        r.append(rFonts)
    rFonts.set(qn('w:eastAsia'), MONO if mono else name)
    rFonts.set(qn('w:ascii'), MONO if mono else name)
    rFonts.set(qn('w:hAnsi'), MONO if mono else name)


def shade_paragraph(paragraph, color_hex):
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color_hex)
    pPr.append(shd)


def add_inline_text(paragraph, text, base_color=None):
    """解析 **加粗** 与 `行内代码`（逐字符状态机，正确处理交错嵌套）。"""
    i, n = 0, len(text)
    bold = False
    code = False
    buf = []

    def flush():
        nonlocal buf
        if not buf:
            return
        s = ''.join(buf)
        run = paragraph.add_run(s)
        if code:
            set_run_font(run, mono=True, size=CODE_SIZE, color=(0xC0, 0x25, 0x25))
        else:
            set_run_font(run, bold=bold, color=base_color)
        buf = []

    while i < n:
        c = text[i]
        if c == '`':
            flush()
            code = not code
            i += 1
        elif text[i:i + 2] == '**' and not code:
            flush()
            bold = not bold
            i += 2
        else:
            buf.append(c)
            i += 1
    flush()


def add_inline_text_styled(paragraph, text, bold=False, color=None, size=BODY_SIZE, name=CN_BODY):
    """标题/引用等整体带样式场景下的行内解析（整体用指定字体，代码仍用等宽）。"""
    code_re = re.compile(r'(`[^`]+`)')
    for seg in code_re.split(text):
        if not seg:
            continue
        if seg.startswith('`') and seg.endswith('`'):
            run = paragraph.add_run(seg[1:-1])
            set_run_font(run, mono=True, size=size, bold=bold, color=(0xC0, 0x25, 0x25))
        else:
            run = paragraph.add_run(seg)
            set_run_font(run, name=name, size=size, bold=bold, color=color)


def convert(md_path, out_path):
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.read().splitlines()

    doc = Document()
    # 页面边距
    for sec in doc.sections:
        sec.left_margin = Inches(0.9)
        sec.right_margin = Inches(0.9)
        sec.top_margin = Inches(0.8)
        sec.bottom_margin = Inches(0.8)

    i = 0
    n = len(lines)
    in_code = False
    code_buf = []

    def flush_code():
        nonlocal code_buf
        if not code_buf:
            return
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.15)
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(6)
        shade_paragraph(p, CODE_BG)
        for ci, cline in enumerate(code_buf):
            if ci > 0:
                run = p.add_run()
                run.add_break()
            run = p.add_run(cline if cline else ' ')
            set_run_font(run, mono=True, size=CODE_SIZE)
        code_buf = []

    while i < n:
        line = lines[i]

        # 代码块
        if line.strip().startswith('```'):
            if not in_code:
                in_code = True
                code_buf = []
            else:
                in_code = False
                flush_code()
            i += 1
            continue
        if in_code:
            code_buf.append(line)
            i += 1
            continue

        stripped = line.strip()

        # 分隔线
        if re.match(r'^-{3,}$', stripped):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            pPr = p._p.get_or_add_pPr()
            pBdr = OxmlElement('w:pBdr')
            bottom = OxmlElement('w:bottom')
            bottom.set(qn('w:val'), 'single')
            bottom.set(qn('w:sz'), '6')
            bottom.set(qn('w:color'), 'AAAAAA')
            pBdr.append(bottom)
            pPr.append(pBdr)
            i += 1
            continue

        # 标题
        m = re.match(r'^(#{1,4})\s+(.*)$', line)
        if m:
            level = len(m.group(1))
            text = m.group(2).strip()
            p = doc.add_paragraph()
            size = {1: H1_SIZE, 2: H2_SIZE, 3: H3_SIZE, 4: H4_SIZE}[level]
            color = {1: (0x1F, 0x38, 0x64), 2: (0x2E, 0x50, 0x86),
                     3: (0x40, 0x40, 0x40), 4: (0x40, 0x40, 0x40)}[level]
            # 标题整体加粗 + 指定字号/颜色，但其中的 `行内代码` 用等宽
            add_inline_text_styled(p, text, bold=True, color=color, size=size, name=CN_HEAD)
            p.paragraph_format.space_before = Pt(14 if level <= 2 else 10)
            p.paragraph_format.space_after = Pt(6 if level <= 2 else 4)
            i += 1
            continue

        # 表格（连续多行以 | 开头）
        if stripped.startswith('|') and i + 1 < n and re.match(r'^\s*\|[\s:\-|]+\|\s*$', lines[i + 1]):
            header = [c.strip() for c in stripped.strip('|').split('|')]
            i += 2  # 跳过分隔行
            rows = [header]
            while i < n and lines[i].strip().startswith('|'):
                rows.append([c.strip() for c in lines[i].strip().strip('|').split('|')])
                i += 1
            ncol = len(header)
            table = doc.add_table(rows=0, cols=ncol)
            table.style = 'Table Grid'
            for ri, row in enumerate(rows):
                cells = table.add_row().cells
                for ci in range(ncol):
                    cell = cells[ci]
                    cell.text = ''
                    cp = cell.paragraphs[0]
                    txt = row[ci] if ci < len(row) else ''
                    # 表头加粗
                    if ri == 0:
                        run = cp.add_run(txt)
                        set_run_font(run, size=BODY_SIZE, bold=True)
                    else:
                        add_inline_text(cp, txt)
            doc.add_paragraph()
            continue

        # 无序列表
        m = re.match(r'^[-*]\s+(.*)$', line)
        if m:
            p = doc.add_paragraph(style='List Bullet')
            add_inline_text(p, m.group(1))
            p.paragraph_format.space_after = Pt(2)
            i += 1
            continue

        # 有序列表
        m = re.match(r'^\d+\.\s+(.*)$', line)
        if m:
            p = doc.add_paragraph(style='List Number')
            add_inline_text(p, m.group(1))
            p.paragraph_format.space_after = Pt(2)
            i += 1
            continue

        # 引用
        m = re.match(r'^>\s?(.*)$', line)
        if m:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.2)
            add_inline_text(p, m.group(1), base_color=(0x60, 0x60, 0x60))
            i += 1
            continue

        # 空行
        if not stripped:
            i += 1
            continue

        # 普通段落
        p = doc.add_paragraph()
        add_inline_text(p, stripped)
        p.paragraph_format.space_after = Pt(4)
        i += 1

    flush_code()

    doc.save(out_path)
    print("OK ->", out_path)


if __name__ == '__main__':
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    src = sys.argv[1]
    dst = sys.argv[2] if len(sys.argv) > 2 else os.path.splitext(src)[0] + '.docx'
    convert(src, dst)
